"""PDF to PPTX conversion using PyMuPDF + python-pptx."""
import asyncio
import io

import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from ..utils.cleanup import open_pdf_document
from ..utils.filenames import temp_output
from ..utils.render import plan_renders, safe_get_pixmap


def _convert_to_pptx_sync(input_path: str) -> str:
    # A PDF that needs a password, or cannot be read, is refused with a 400.
    doc = open_pdf_document(input_path)
    try:
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        # 200 DPI, or less for a page too large for the pixel cap; the whole
        # request held to its render budget first.
        zooms = plan_renders(doc, 200 / 72, advice="Split the PDF and convert the parts separately.")

        for page_num in range(len(doc)):
            page = doc[page_num]
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

            # Render at 200 DPI for crisp slides on modern displays.
            zoom = zooms[page_num]
            pix = safe_get_pixmap(page, matrix=fitz.Matrix(zoom, zoom))
            width, height = pix.width, pix.height

            # JPEG for smaller files (slides don't need transparency).
            img_data = io.BytesIO()
            with Image.frombytes("RGB", [width, height], pix.samples_mv) as img:
                del pix  # PIL has its own copy
                img.save(img_data, format="JPEG", quality=90, optimize=True)
            img_data.seek(0)

            # Calculate dimensions to fit slide
            slide_w = prs.slide_width
            slide_h = prs.slide_height
            img_ratio = width / height
            slide_ratio = slide_w / slide_h

            if img_ratio > slide_ratio:
                width = slide_w
                height = int(slide_w / img_ratio)
            else:
                height = slide_h
                width = int(slide_h * img_ratio)

            left = int((slide_w - width) / 2)
            top = int((slide_h - height) / 2)

            slide.shapes.add_picture(img_data, left, top, width, height)
    finally:
        doc.close()

    output_path = temp_output("converted", "pptx")
    prs.save(str(output_path))
    return str(output_path)


async def convert_to_pptx(input_path: str) -> str:
    # The body is pure CPU/PyMuPDF work; offload it so it doesn't block the
    # event loop (this was declared async but ran synchronously before).
    return await asyncio.to_thread(_convert_to_pptx_sync, input_path)
