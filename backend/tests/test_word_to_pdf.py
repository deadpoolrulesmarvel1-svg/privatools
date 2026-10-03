"""Word to PDF sets the text of each body paragraph in Helvetica.

It reads that text with python-docx, not LibreOffice. Word keeps an equation
(OMML, an m:oMath element) beside the text runs of its paragraph, so the text
python-docx gives leaves it out, and so does the PDF: "Before x/y after." comes
out as "Before after.". The guide has to say so. All documents are synthetic.
"""
from __future__ import annotations

import io

import docx
import fitz
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from backend.app.services import word_to_pdf_service
from backend.app.tool_content import TOOL_FAQ, TOOL_HOWTO

FRACTION = (
    f'<m:oMath {nsdecls("m")}><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num>'
    '<m:den><m:r><m:t>y</m:t></m:r></m:den></m:f></m:oMath>'
)


def word_document(*paragraphs: tuple[str, bool, str]) -> bytes:
    """A .docx with one paragraph per (before, has an x/y equation, after)."""
    document = docx.Document()
    for before, equation, after in paragraphs:
        paragraph = document.add_paragraph(before)
        if equation:
            paragraph._p.append(parse_xml(FRACTION))
        paragraph.add_run(after)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pdf_text(path: str) -> str:
    with fitz.open(path) as pdf:
        return " ".join(page.get_text() for page in pdf).strip()


def converted(tmp_path, content: bytes) -> str:
    source = tmp_path / "equation.docx"
    source.write_bytes(content)
    return word_to_pdf_service.word_to_pdf(str(source))


def test_an_equation_is_left_out_and_the_guide_says_so(tmp_path):
    text = pdf_text(converted(tmp_path, word_document(("Before ", True, " after."))))
    assert text == "Before after."

    steps = " ".join(step["text"] for step in TOOL_HOWTO["word-to-pdf"])
    assert "equations" in steps
    faq = {item["q"]: item["a"] for item in TOOL_FAQ["word-to-pdf"]}
    assert "equations" in faq["Will the PDF look exactly like my Word document?"]
    [answer] = [a for q, a in faq.items() if "equation" in q.lower()]
    assert "Before after." in answer and "Office to PDF" in answer
