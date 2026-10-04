"""Table detection leaves PyMuPDF's process-wide glyph-height switch off.

PyMuPDF 1.28's Page.find_tables turns TOOLS.set_small_glyph_heights on, then
reads the page's rotation, and only after that enters the try block whose
finally turns the switch back. On a page whose rotation it cannot read, as in a
damaged PDF, the switch stayed on for every later text extraction in the
process: word boxes came out one font size tall, and on text turned 180
degrees, beside the text, which is what Redact covers. CI saw it as two failures
of test_page_contract after test_damaged_pdfs_everywhere had sent cut-short
PDFs to Extract Tables and PDF to Excel (on x86_64; the dev VM's ARM build read
those pages without failing).

The tests make the rotation step fail the way a damaged page does, in
PyMuPDF's own find_tables.
"""

from __future__ import annotations

import fitz  # PyMuPDF
import pymupdf.table
import pytest

from backend.app.services import pdf_to_excel_service, table_extractor_service


@pytest.fixture
def turned_pdf(tmp_path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "A table that is not one", fontsize=12)
    page.set_rotation(180)
    path = tmp_path / "turned.pdf"
    doc.save(path)
    doc.close()
    return str(path)


@pytest.fixture
def rotation_fails(monkeypatch):
    """find_tables meets a page whose rotation cannot be read, after it has
    turned the switch on and before its try block."""
    def fails(page):
        raise RuntimeError("code=7: Invalid number of pages")

    monkeypatch.setattr(pymupdf.table, "page_rotation_set0", fails)
    fitz.TOOLS.set_small_glyph_heights(False)
    yield
    fitz.TOOLS.set_small_glyph_heights(False)


def test_the_failure_is_one_pymupdf_leaves_the_switch_on_after(turned_pdf, rotation_fails):
    with fitz.open(turned_pdf) as doc:
        with pytest.raises(RuntimeError):
            doc[0].find_tables()
    assert fitz.TOOLS.set_small_glyph_heights(), "PyMuPDF no longer leaks it: the wrapper may go"


def test_extract_tables_leaves_the_switch_off(turned_pdf, rotation_fails):
    with pytest.raises(Exception):
        table_extractor_service.extract_tables(turned_pdf)
    assert not fitz.TOOLS.set_small_glyph_heights()


def test_pdf_to_excel_leaves_the_switch_off(turned_pdf, rotation_fails):
    # PDF to Excel falls back to the page's text when table detection fails.
    try:
        pdf_to_excel_service._build_workbook(turned_pdf)
    except Exception:  # noqa: BLE001 - only the switch is asked about here
        pass
    assert not fitz.TOOLS.set_small_glyph_heights()


def _ruled_tables(path, pages: int) -> None:
    doc = fitz.open()
    for p in range(pages):
        page = doc.new_page()
        rows = 12 + p % 7
        for row in range(rows):
            for col in range(5):
                page.insert_text((62 + 100 * col, 80 + 18 * row), f"p{p}r{row}c{col} g", fontsize=9)
        for y in range(66, 66 + 18 * (rows + 1), 18):
            page.draw_line((58, y), (558, y), width=0.5)
        for x in range(58, 559, 100):
            page.draw_line((x, 66), (x, 66 + 18 * rows), width=0.5)
    doc.save(path)
    doc.close()


def _cells(path) -> list:
    from openpyxl import load_workbook

    out, _ = pdf_to_excel_service._build_workbook(str(path))
    try:
        return [(ws.title, c.coordinate, c.value) for ws in load_workbook(out) for row in ws.iter_rows()
                for c in row if c.value]
    finally:
        import os

        os.remove(out)


def test_pdf_to_excel_reads_pages_in_parallel_as_it_reads_them_one_by_one(tmp_path, monkeypatch):
    # On a host with more than two CPUs the service runs find_tables on up to
    # four pages at once. A call that ended turned the shared switch off under
    # the others, and their cells lost the spaces between words ("p9r0c0 g"
    # came out "p9r0c0g") in most runs.
    import os

    path = tmp_path / "tables.pdf"
    _ruled_tables(path, 12)
    monkeypatch.setattr(os, "cpu_count", lambda: 1)
    one_by_one = _cells(path)
    monkeypatch.setattr(os, "cpu_count", lambda: 4)
    for _ in range(3):
        assert _cells(path) == one_by_one
    assert not fitz.TOOLS.set_small_glyph_heights()


def test_smart_redact_never_searches_while_a_table_detection_holds_the_switch(tmp_path, monkeypatch):
    import threading

    from backend.app.services import smart_redact_service
    from backend.app.utils import tables

    doc = fitz.open()
    doc.new_page().insert_text((300, 300), "SECRET 4111-1111-1111-1111", fontsize=14, rotate=180)
    path = tmp_path / "secret.pdf"
    doc.save(path)
    doc.close()

    holding, done = threading.Event(), threading.Event()

    def slow_find_tables(self, **kwargs):  # PyMuPDF's switch, held for a while
        fitz.TOOLS.set_small_glyph_heights(True)
        holding.set()
        done.wait(1.0)  # until Smart Redact has searched, or a second at most

    seen = []
    search_for = fitz.Page.search_for

    def recording_search_for(self, *args, **kwargs):
        seen.append(fitz.TOOLS.set_small_glyph_heights())
        done.set()
        return search_for(self, *args, **kwargs)

    monkeypatch.setattr(fitz.Page, "find_tables", slow_find_tables)
    monkeypatch.setattr(fitz.Page, "search_for", recording_search_for)
    with fitz.open() as blank:
        detection = threading.Thread(target=tables.find_tables, args=(blank.new_page(),))
        detection.start()
        assert holding.wait(5)
        out, hits = smart_redact_service.smart_redact(str(path), ["4111-1111-1111-1111"])
        detection.join(5)
    assert hits == 1
    assert seen == [False]
    with fitz.open(out) as redacted:
        assert "4111" not in redacted[0].get_text()
    assert not fitz.TOOLS.set_small_glyph_heights()


def test_a_table_is_still_found(tmp_path):
    doc = fitz.open()
    page = doc.new_page()
    for row in range(4):
        for col in range(3):
            page.insert_text((80 + 120 * col, 118 + 30 * row), f"r{row}c{col}", fontsize=11)
    for y in range(100, 221, 30):
        page.draw_line((70, y), (430, y))
    for x in (70, 190, 310, 430):
        page.draw_line((x, 100), (x, 220))
    path = tmp_path / "table.pdf"
    doc.save(path)
    doc.close()
    out = table_extractor_service.extract_tables(str(path))
    with open(out, encoding="utf-8") as fh:
        assert "r3c2" in fh.read()
    assert not fitz.TOOLS.set_small_glyph_heights()
